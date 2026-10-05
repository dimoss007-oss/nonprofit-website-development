import base64
import copy
import io
import json
import os
from datetime import date, datetime

import psycopg2
from docx import Document
from template_data import TEMPLATE_B64
from psycopg2.extras import RealDictCursor

SCHEMA = os.environ.get("MAIN_DB_SCHEMA", "public")
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

CORS = {
    "Access-Control-Allow-Origin": "*",
    "Access-Control-Allow-Methods": "GET, POST, OPTIONS",
    "Access-Control-Allow-Headers": "Content-Type",
}

MONTHS = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]


def ok(data, status=200):
    return {"statusCode": status, "headers": {**CORS, "Content-Type": "application/json"}, "body": json.dumps(data, default=str, ensure_ascii=False)}


def err(msg, status=400):
    return {"statusCode": status, "headers": {**CORS, "Content-Type": "application/json"}, "body": json.dumps({"error": msg}, ensure_ascii=False)}


def set_runs(par, mapping: dict):
    """mapping: индекс run -> новый текст. Форматирование runs сохраняется."""
    for idx, text in mapping.items():
        par.runs[idx].text = text


def find_par(doc, prefix: str, start_after: int = 0):
    for i, p in enumerate(doc.paragraphs):
        if i >= start_after and p.text.strip().startswith(prefix):
            return i, p
    raise ValueError(f"В шаблоне не найден абзац: {prefix}")


def fmt_date(d) -> str:
    return f"{d.day:02d}.{d.month:02d}.{d.year}"


def child_relation(middle_name: str) -> str:
    m = (middle_name or "").strip().lower()
    if m.endswith("на"):
        return "дочь"
    if m.endswith("ич"):
        return "сын"
    return "ребёнок"


def build_contract(patient: dict, children: list, number: str, contract_date: date) -> bytes:
    doc = Document(io.BytesIO(base64.b64decode(TEMPLATE_B64)))

    last = (patient.get("last_name") or "").strip()
    first = (patient.get("first_name") or "").strip()
    middle = (patient.get("middle_name") or "").strip()
    full_name = " ".join(x for x in (last, first, middle) if x)
    short_name = last
    initials = "".join(f"{x[0]}. " for x in (first, middle) if x).strip()
    if initials:
        short_name = f"{last} {initials}"

    birth = patient.get("birth_date")
    birth_year = str(birth.year) if birth else "____"
    series = (patient.get("passport_series") or "").strip() or "____"
    pnumber = (patient.get("passport_number") or "").strip() or "______"
    issued_by = (patient.get("passport_issued_by") or "").strip()
    issued_date = patient.get("passport_issued_date")
    issued_text = issued_by
    if issued_date:
        issued_text = f"{issued_by}, {fmt_date(issued_date)}".strip(", ")
    issued_text = issued_text or "______________________"
    address = (patient.get("address") or "").strip() or "______________________"

    day = f"{contract_date.day:02d}"
    month = MONTHS[contract_date.month - 1]
    year = str(contract_date.year)

    _, p = find_par(doc, "ДОГОВОР №")
    set_runs(p, {1: number, 2: "", 3: "", 4: "", 5: ""})

    _, p = find_par(doc, "г. Пенза")
    set_runs(p, {6: "«", 7: day, 8: f"» {month} ", 9: year, 10: " г.", 11: "", 12: "", 13: ""})

    _, p = find_par(doc, "АВТОНОМНАЯ НЕКОММЕРЧЕСКАЯ")
    set_runs(p, {4: full_name, 5: "", 7: birth_year})

    _, p = find_par(doc, "5.1.")
    set_runs(p, {1: "«", 2: day, 3: "» ", 4: month, 5: "", 6: " ", 7: year, 8: ""})

    _, p = find_par(doc, "Я,", 140)
    set_runs(p, {1: full_name, 2: ""})

    idx_family, _ = find_par(doc, "Совместно с")
    child_par = None
    for j in range(idx_family + 1, len(doc.paragraphs)):
        if doc.paragraphs[j].text.strip():
            child_par = doc.paragraphs[j]
            break
    if child_par is None:
        raise ValueError("В шаблоне не найден абзац со списком детей")

    lines = []
    for c in children:
        c_name = " ".join(x for x in ((c.get("last_name") or "").strip(), (c.get("first_name") or "").strip(), (c.get("middle_name") or "").strip()) if x)
        c_birth = fmt_date(c["birth_date"]) if c.get("birth_date") else "__.__.____"
        lines.append((c_name, c_birth, child_relation(c.get("middle_name"))))
    if not lines:
        lines = [("—", "", "")]

    anchor = child_par._p
    for n, (c_name, c_birth, rel) in enumerate(lines):
        if n == 0:
            target = child_par
        else:
            new_el = copy.deepcopy(child_par._p)
            anchor.addnext(new_el)
            anchor = new_el
            from docx.text.paragraph import Paragraph
            target = Paragraph(new_el, child_par._parent)
        text_rel = f"– {rel}" if rel else ""
        set_runs(target, {0: c_name + " ", 1: "", 2: "", 3: c_birth, 4: " " if rel else "", 5: text_rel})
        for k in range(6, len(target.runs)):
            target.runs[k].text = ""

    _, p = find_par(doc, "Я,", idx_family)
    set_runs(p, {2: full_name, 3: ""})

    _, p = find_par(doc, "основной документ")
    set_runs(p, {6: f". {series}", 9: pnumber})

    _, p = find_par(doc, "Адрес регистрации", 150)
    set_runs(p, {3: address, 4: "", 5: "", 6: "", 7: ""})

    for par in doc.paragraphs[190:]:
        if par.text.strip().startswith("«") and "/____" in par.text:
            set_runs(par, {0: f"«{day}» {month} {year} г.", 1: "", 2: "", 3: "", 4: "", 5: "", 6: "", 7: ""})
            break

    cell = doc.tables[0].rows[0].cells[3]
    for par in cell.paragraphs:
        t = par.text.strip()
        if t.startswith("ФИО:"):
            set_runs(par, {1: short_name, 2: ""})
        elif t.startswith("Паспорт:"):
            set_runs(par, {4: series, 5: "", 8: pnumber, 9: ""})
        elif t.startswith("Выдан:"):
            set_runs(par, {1: issued_text})
        elif t.startswith("Адрес регистрации"):
            set_runs(par, {3: address, 4: "", 5: "", 6: "", 7: ""})

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def next_number(cur, contract_date: date) -> str:
    cur.execute(
        f"SELECT contract_number FROM {SCHEMA}.contracts WHERE date_trunc('month', contract_date) = date_trunc('month', %s::date)",
        (contract_date.isoformat(),),
    )
    month = f"{contract_date.month:02d}"
    biggest = 0
    for row in cur.fetchall():
        head, _, tail = str(row["contract_number"]).partition("/")
        if head.strip().isdigit() and tail.strip() == month:
            biggest = max(biggest, int(head))
    return f"{biggest + 1:02d}/{month}"


def parse_date(raw):
    return datetime.strptime(raw, "%Y-%m-%d").date() if raw else date.today()


def handler(event: dict, context) -> dict:
    """Договоры о предоставлении комплекса социальных услуг (.docx) по данным пациента из CRM.
    GET ?contract_date=YYYY-MM-DD -> {next_number} следующий номер вида NN/MM в рамках месяца.
    POST {patient_id, contract_number, contract_date (YYYY-MM-DD, необязательно)} -> {file_name, file_base64, content_type}
    и запоминает выданный номер."""
    if event.get("httpMethod") == "OPTIONS":
        return {"statusCode": 200, "headers": CORS, "body": ""}

    if event.get("httpMethod") == "GET":
        params = event.get("queryStringParameters") or {}
        try:
            d = parse_date(params.get("contract_date"))
        except ValueError:
            return err("Неверная дата договора")
        conn = psycopg2.connect(os.environ["DATABASE_URL"])
        try:
            cur = conn.cursor(cursor_factory=RealDictCursor)
            number = next_number(cur, d)
            cur.close()
        finally:
            conn.close()
        return ok({"next_number": number})

    body = json.loads(event.get("body") or "{}")
    patient_id = body.get("patient_id")
    number = str(body.get("contract_number") or "").strip()
    if not patient_id:
        return err("Поле patient_id обязательно")
    if not number:
        return err("Укажите номер договора")

    raw_date = body.get("contract_date")
    try:
        contract_date = parse_date(raw_date)
    except ValueError:
        return err("Неверная дата договора")

    conn = psycopg2.connect(os.environ["DATABASE_URL"])
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(f"SELECT * FROM {SCHEMA}.patients WHERE id = %s", (int(patient_id),))
        patient = cur.fetchone()
        if not patient:
            return err("Пациент не найден", 404)
        cur.execute(f"SELECT * FROM {SCHEMA}.patient_children WHERE patient_id = %s ORDER BY birth_date", (int(patient_id),))
        children = cur.fetchall()
        cur.execute(
            f"SELECT 1 FROM {SCHEMA}.contracts WHERE patient_id = %s AND contract_number = %s AND contract_date = %s",
            (int(patient_id), number, contract_date.isoformat()),
        )
        if not cur.fetchone():
            cur.execute(
                f"INSERT INTO {SCHEMA}.contracts (patient_id, contract_number, contract_date) VALUES (%s, %s, %s)",
                (int(patient_id), number, contract_date.isoformat()),
            )
            conn.commit()
        cur.close()
    finally:
        conn.close()

    data = build_contract(dict(patient), [dict(c) for c in children], number, contract_date)
    safe_number = number.replace("/", "-").replace(" ", "_")
    last = (patient.get("last_name") or "Пациент").replace(" ", "_")
    return ok({
        "file_name": f"Договор_{safe_number}_{last}.docx",
        "file_base64": base64.b64encode(data).decode("ascii"),
        "content_type": DOCX_TYPE,
    })
