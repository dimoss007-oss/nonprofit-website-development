import base64
from crm_common import SCHEMA, err, is_admin, ok
from local_summary import generate_text_summary
from yandex_ai import build_characteristic_docx, generate_official_characteristic, generate_yandex_summary, get_system_prompt


def get_text_summary(cur, conn, params, body):
    patient_id = params.get("id")
    try:
        days = int(params.get("days", 7))
    except (TypeError, ValueError):
        days = 7
    days = max(1, min(days, 90))
    result = generate_text_summary(cur, patient_id, SCHEMA, days)
    return ok(result)


def get_yandex_summary(cur, conn, params, body):
    patient_id = params.get("id")
    try:
        days = int(params.get("days", 7))
    except (TypeError, ValueError):
        days = 7
    days = max(1, min(days, 90))
    result = generate_yandex_summary(cur, patient_id, SCHEMA, days)
    return ok(result)


def get_ai_settings(cur, conn, params, body):
    system_prompt = get_system_prompt(cur, SCHEMA)
    return ok({"yandexgpt_system_prompt": system_prompt})


def update_ai_settings(cur, conn, params, body):
    auth_login = body.get("auth_login", "")
    auth_password = body.get("auth_password", "")
    if not is_admin(cur, SCHEMA, auth_login, auth_password):
        return err("Нет прав: изменение промпта доступно только администраторам", 403)

    new_prompt = (body.get("yandexgpt_system_prompt") or "").strip()
    if not new_prompt:
        return err("Текст системного промпта не может быть пустым")

    cur.execute(
        f"""INSERT INTO {SCHEMA}.crm_settings (id, yandexgpt_system_prompt, updated_by, updated_at)
            VALUES (1, %s, %s, NOW())
            ON CONFLICT (id) DO UPDATE SET yandexgpt_system_prompt = EXCLUDED.yandexgpt_system_prompt,
                updated_by = EXCLUDED.updated_by, updated_at = NOW()
            RETURNING yandexgpt_system_prompt, updated_at, updated_by""",
        (new_prompt, auth_login)
    )
    settings = cur.fetchone()
    conn.commit()
    return ok({"settings": dict(settings)})


def generate_and_save_yandex_summary(cur, conn, params, body):
    pid = body.get("patient_id")
    if not pid:
        return err("Поле patient_id обязательно")
    try:
        days = int(body.get("days", 3))
    except (TypeError, ValueError):
        days = 3
    days = max(1, min(days, 90))

    result = generate_yandex_summary(cur, pid, SCHEMA, days)
    summary_text = result.get("summary_text")
    if not summary_text:
        return err(result.get("error") or "Не удалось сформировать сводку", 502)

    cur.execute(
        f"INSERT INTO {SCHEMA}.patient_ai_summaries (patient_id, summary_text, source) VALUES (%s, %s, 'yandex_gpt') RETURNING id, patient_id, summary_text, source, created_at",
        (pid, summary_text)
    )
    summary = cur.fetchone()
    conn.commit()
    return ok({"summary": dict(summary)}, 201)


def generate_characteristic_docx(cur, conn, params, body):
    pid = body.get("patient_id")
    if not pid:
        return err("Поле patient_id обязательно")

    result = generate_official_characteristic(cur, pid, SCHEMA)
    if result.get("error"):
        return err(result["error"], 502)

    docx_bytes = build_characteristic_docx(result)
    file_b64 = base64.b64encode(docx_bytes).decode("ascii")
    file_name = f"Характеристика_{result['full_name'].replace(' ', '_')}.docx"

    return ok({
        "file_name": file_name,
        "file_base64": file_b64,
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    })


def save_local_summary(cur, conn, params, body):
    pid = body.get("patient_id")
    summary_text = (body.get("summary_text") or "").strip()
    source = body.get("source") or "rule_based"
    if source not in ("rule_based", "yandex_gpt"):
        source = "rule_based"
    if not pid or not summary_text:
        return err("Поля patient_id и summary_text обязательны")

    cur.execute(
        f"INSERT INTO {SCHEMA}.patient_ai_summaries (patient_id, summary_text, source) VALUES (%s, %s, %s) RETURNING id, patient_id, summary_text, source, created_at",
        (pid, summary_text, source)
    )
    summary = cur.fetchone()
    conn.commit()
    return ok({"summary": dict(summary)}, 201)


def get_psychologist_reports(cur, conn, params, body):
    patient_id = params.get("id")
    cur.execute(
        f"""SELECT id, author, report_date, report_text, created_at
            FROM {SCHEMA}.psychologist_reports
            WHERE patient_id = %s ORDER BY report_date DESC, id DESC""",
        (patient_id,),
    )
    return ok({"reports": [dict(r) for r in cur.fetchall()]})


def _psych_day_rows(cur, report_date):
    cur.execute(
        f"""SELECT r.id, r.patient_id, r.author, r.report_date, r.report_text,
                   p.first_name, p.last_name, p.alias
            FROM {SCHEMA}.psychologist_reports r
            JOIN {SCHEMA}.patients p ON p.id = r.patient_id
            WHERE r.report_date = %s
            ORDER BY p.last_name, p.first_name, r.id""",
        (report_date,),
    )
    rows = [dict(r) for r in cur.fetchall()]
    for r in rows:
        r["patient_name"] = " ".join(x for x in (r["last_name"], r["first_name"]) if x)
        if r.get("alias"):
            r["patient_name"] += f" ({r['alias']})"
    return rows


def get_psychologist_day(cur, conn, params, body):
    report_date = (params.get("date") or "").strip()
    if not report_date:
        return err("Укажите дату")
    return ok({"reports": _psych_day_rows(cur, report_date)})


def generate_psychologist_day_docx(cur, conn, params, body):
    import io
    from datetime import datetime
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    report_date = (body.get("date") or "").strip()
    if not report_date:
        return err("Укажите дату")
    rows = _psych_day_rows(cur, report_date)
    if not rows:
        return err("За эту дату нет отчётов психолога")
    d = datetime.strptime(report_date[:10], "%Y-%m-%d").strftime("%d.%m.%Y")

    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    title = doc.add_heading("Отчёт психолога по пациентам", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(f"Дата: {d}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors = ", ".join(dict.fromkeys(r["author"] for r in rows if r.get("author")))
    if authors:
        doc.add_paragraph(f"Автор: {authors}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in rows:
        doc.add_heading(r["patient_name"], level=2)
        for line in (r["report_text"] or "").split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return ok({
        "file_name": f"Отчёт психолога {report_date}.docx",
        "file_base64": base64.b64encode(buf.getvalue()).decode("ascii"),
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    })
