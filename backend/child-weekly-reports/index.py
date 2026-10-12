import difflib
import base64
import io
import json
import os
import re
from datetime import datetime, timedelta

import psycopg2
import requests
from psycopg2.extras import RealDictCursor

SCHEMA = os.environ.get("MAIN_DB_SCHEMA", "public")

CORS = {
    "Access-Control-Allow-Origin": "*",
    "Access-Control-Allow-Methods": "GET, POST, PUT, DELETE, OPTIONS",
    "Access-Control-Allow-Headers": "Content-Type, X-User-Id, X-Auth-Token",
}


def get_conn():
    return psycopg2.connect(os.environ["DATABASE_URL"])


def ok(data, status=200):
    return {"statusCode": status, "headers": {**CORS, "Content-Type": "application/json"}, "body": json.dumps(data, default=str, ensure_ascii=False)}


def err(msg, status=400):
    return {"statusCode": status, "headers": {**CORS, "Content-Type": "application/json"}, "body": json.dumps({"error": msg}, ensure_ascii=False)}

LINE_RE = re.compile(r"^\s*(?:\d+\s*[.)]\s*)?([^\-\u2014\u2013:]{2,40}?)\s*[-\u2014\u2013:]\s*(.+)$")
NAME_CUTOFF = 0.8

DIMINUTIVES = {
    "аня": "анна", "катя": "екатерина", "света": "светлана", "настя": "анастасия",
    "таня": "татьяна", "наташа": "наталья", "саша": "александр", "лена": "елена",
    "оля": "ольга", "женя": "евгений", "дима": "дмитрий", "маша": "мария",
    "юля": "юлия", "надя": "надежда", "люда": "людмила", "вика": "виктория",
    "ксюша": "ксения", "лиза": "елизавета", "даша": "дарья", "паша": "павел",
    "коля": "николай", "миша": "михаил", "вова": "владимир", "гена": "геннадий",
    "толя": "анатолий", "валя": "валентина", "галя": "галина", "тоня": "антонина",
    "стас": "станислав", "сережа": "сергей", "серёжа": "сергей", "мирося": "мирослава",
}


def norm(value: str) -> str:
    return re.sub(r"[.\s]+", " ", (value or "").strip().lower().replace("ё", "е")).strip()


def monday_of(raw: str) -> str:
    d = datetime.strptime(raw[:10], "%Y-%m-%d").date()
    return (d - timedelta(days=d.weekday())).isoformat()


def load_children(cur) -> list:
    cur.execute(
        f"""SELECT c.id, c.first_name, c.last_name, c.middle_name, c.alias,
                   p.first_name AS mother_first_name, p.last_name AS mother_last_name, p.alias AS mother_alias
            FROM {SCHEMA}.patient_children c
            JOIN {SCHEMA}.patients p ON p.id = c.patient_id
            WHERE p.discharge_date IS NULL AND (p.care_stage IS NULL OR p.care_stage = 'inpatient')
            ORDER BY c.last_name, c.first_name"""
    )
    rows = [dict(r) for r in cur.fetchall()]
    for r in rows:
        r["display_name"] = " ".join(x for x in (r["last_name"], r["first_name"]) if x)
        if r.get("alias"):
            r["display_name"] += f" ({r['alias']})"
    return rows


def build_index(children: list) -> dict:
    index = {}
    for c in children:
        first = norm(c["first_name"])
        last = norm(c["last_name"])
        keys = {first}
        alias = norm(c.get("alias") or "")
        if alias:
            keys.add(alias)
        if last:
            keys.update({last, f"{first} {last}", f"{last} {first}", f"{first} {last[:1]}"})
        for key in keys:
            if key:
                index.setdefault(key, []).append(c["id"])
    return index


def expand(name: str) -> str:
    parts = name.split(" ", 1)
    full = DIMINUTIVES.get(parts[0])
    if not full:
        return name
    return f"{full} {parts[1] if len(parts) > 1 else ''}".strip()


def match_children(name: str, index: dict) -> list:
    n = norm(name)
    if not n:
        return []
    variants = [n]
    e = norm(expand(n))
    if e != n:
        variants.append(e)
    for v in variants:
        if v in index:
            return list(dict.fromkeys(index[v]))
    for v in variants:
        close = difflib.get_close_matches(v, list(index.keys()), n=1, cutoff=NAME_CUTOFF)
        if close:
            return list(dict.fromkeys(index[close[0]]))
    return []


def parse_report(text: str, children: list) -> list:
    """Разбирает текст построчно: строка вида «Имя - состояние» начинает блок ребёнка,
    строки без имени дописываются к текущему блоку."""
    index = build_index(children)
    by_id = {c["id"]: c for c in children}
    blocks = []
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        m = LINE_RE.match(line)
        name_part = m.group(1).strip() if m else ""
        words = name_part.split()
        found = match_children(name_part, index) if m and 1 <= len(words) <= 3 else []
        looks_like_name = bool(m) and 1 <= len(words) <= 3 and (bool(found) or name_part[:1].isupper())
        if looks_like_name:
            blocks.append({"name_in_text": name_part, "text": m.group(2).strip(), "candidate_ids": found})
        elif blocks:
            blocks[-1]["text"] += "\n" + line
    result = []
    for b in blocks:
        ids = b["candidate_ids"]
        result.append({
            "name_in_text": b["name_in_text"],
            "text": b["text"],
            "child_id": ids[0] if len(ids) == 1 else None,
            "child_name": by_id[ids[0]]["display_name"] if len(ids) == 1 else None,
            "status": "matched" if len(ids) == 1 else ("ambiguous" if len(ids) > 1 else "unmatched"),
            "candidates": [{"id": i, "name": by_id[i]["display_name"], "mother": by_id[i]["mother_alias"] or by_id[i]["mother_last_name"]} for i in ids],
        })
    return result


YANDEX_GPT_URL = "https://llm.api.cloud.yandex.net/foundationModels/v1/completion"
SUMMARY_PROMPT = (
    "Ты помощник психолога центра помощи матерям с детьми. Тебе дают еженедельные отчёты сотрудников по детям. "
    "Составь краткую выжимку на русском языке: 1) общая картина за неделю в 2-3 предложениях; "
    "2) по каждому ребёнку одна строка: имя — главное (состояние, поведение, успехи); "
    "3) отдельный список «Требуют внимания» — только если в отчётах есть тревожные сигналы. "
    "Ничего не выдумывай, опирайся только на текст отчётов. Без медицинских диагнозов."
)
MAX_SUMMARY_INPUT = 12000


def load_week_reports(cur, week_start: str) -> list:
    cur.execute(
        f"""SELECT r.id, r.child_id, r.author, r.week_start, r.report_text, r.updated_at,
                   c.first_name, c.last_name, c.alias
            FROM {SCHEMA}.child_weekly_reports r
            JOIN {SCHEMA}.patient_children c ON c.id = r.child_id
            WHERE r.week_start = %s
            ORDER BY c.last_name, c.first_name, r.id""",
        (week_start,),
    )
    rows = [dict(r) for r in cur.fetchall()]
    for r in rows:
        r["child_name"] = " ".join(x for x in (r["last_name"], r["first_name"]) if x)
        if r.get("alias"):
            r["child_name"] += f" ({r['alias']})"
    return rows


def generate_summary(rows: list) -> str:
    folder_id = (os.environ.get("YANDEX_FOLDER_ID") or "").strip()
    api_key = (os.environ.get("YANDEX_API_KEY") or "").strip()
    if not folder_id or not api_key:
        raise RuntimeError("Не настроен доступ к YandexGPT")
    chunks, total = [], 0
    for r in rows:
        chunk = f"{r['child_name']}:\n{r['report_text'].strip()}"
        if total + len(chunk) > MAX_SUMMARY_INPUT:
            break
        chunks.append(chunk)
        total += len(chunk)
    payload = {
        "modelUri": f"gpt://{folder_id}/yandexgpt/latest",
        "completionOptions": {"stream": False, "temperature": 0.3, "maxTokens": 1200},
        "messages": [
            {"role": "system", "text": SUMMARY_PROMPT},
            {"role": "user", "text": "\n\n".join(chunks)},
        ],
    }
    resp = requests.post(YANDEX_GPT_URL, headers={"Authorization": f"Api-Key {api_key}", "Content-Type": "application/json"}, json=payload, timeout=60)
    if resp.status_code != 200:
        raise RuntimeError("YandexGPT вернул ошибку, попробуйте позже")
    return resp.json()["result"]["alternatives"][0]["message"]["text"].strip()


def week_range_label(week_start: str) -> str:
    d = datetime.strptime(week_start[:10], "%Y-%m-%d").date()
    e = d + timedelta(days=6)
    return f"{d.strftime('%d.%m.%Y')} – {e.strftime('%d.%m.%Y')}"


def build_summary_docx(week_start: str, summary_text: str, rows: list) -> bytes:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title = doc.add_heading("Выжимка из отчётов по детям", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    period = doc.add_paragraph(f"Неделя: {week_range_label(week_start)}")
    period.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Выжимка", level=2)
    for line in summary_text.replace("**", "").split("\n"):
        if line.strip():
            doc.add_paragraph(line.strip())

    names = list(dict.fromkeys(r["child_name"] for r in rows))
    doc.add_heading(f"Дети, по которым есть отчёты ({len(names)})", level=2)
    for n in names:
        doc.add_paragraph(n, style="List Number")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def handler(event: dict, context) -> dict:
    """Еженедельные отчёты по детям. GET ?child_id=N — список; POST — создать; PUT ?id=N — изменить; DELETE ?id=N — удалить.
    POST {action: parse} — разбор общего текста по детям; POST {action: import} — сохранение разобранного отчёта."""
    if event.get("httpMethod") == "OPTIONS":
        return {"statusCode": 200, "headers": CORS, "body": ""}

    method = event.get("httpMethod", "GET")
    params = event.get("queryStringParameters") or {}
    body = json.loads(event.get("body") or "{}") if method in ("POST", "PUT") else {}

    conn = get_conn()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        if method == "POST" and body.get("action") == "summary":
            if not body.get("week_start"):
                return err("Укажите неделю отчёта")
            rows = load_week_reports(cur, monday_of(body["week_start"]))
            if body.get("child_ids"):
                wanted = {int(i) for i in body["child_ids"]}
                rows = [r for r in rows if r["child_id"] in wanted]
            if not rows:
                return err("За эту неделю нет загруженных отчётов")
            try:
                text = generate_summary(rows)
            except (RuntimeError, requests.exceptions.RequestException) as e:
                return err(str(e) if isinstance(e, RuntimeError) else "Нет связи с YandexGPT, попробуйте позже", 502)
            return ok({"summary_text": text, "reports_count": len(rows)})

        if method == "POST" and body.get("action") == "summary_docx":
            if not body.get("week_start"):
                return err("Укажите неделю отчёта")
            summary_text = (body.get("summary_text") or "").strip()
            if not summary_text:
                return err("Сначала составьте выжимку")
            week = monday_of(body["week_start"])
            rows = load_week_reports(cur, week)
            data = build_summary_docx(week, summary_text, rows)
            return ok({
                "file_name": f"Выжимка по детям {week}.docx",
                "file_base64": base64.b64encode(data).decode("ascii"),
                "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            })

        if method == "POST" and body.get("action") in ("parse", "import"):
            action = body["action"]
            week_start = body.get("week_start")
            if not week_start:
                return err("Укажите неделю отчёта")
            week_start = monday_of(week_start)

            if action == "parse":
                text = (body.get("report_text") or "").strip()
                if not text:
                    return err("Вставьте текст отчёта")
                children = load_children(cur)
                items = parse_report(text, children)
                if not items:
                    return err("Не нашёл ни одной строки вида «Имя - состояние». Проверьте формат отчёта")
                return ok({
                    "week_start": week_start,
                    "items": items,
                    "children": [{"id": c["id"], "name": c["display_name"], "mother": c["mother_alias"] or c["mother_last_name"]} for c in children],
                })

            items = body.get("items") or []
            author = (body.get("author") or "").strip() or None
            merged = {}
            for it in items:
                child_id = it.get("child_id")
                text = (it.get("text") or "").strip()
                if child_id and text:
                    merged[int(child_id)] = (merged.get(int(child_id)) + "\n" if int(child_id) in merged else "") + text
            if not merged:
                return err("Нет отчётов для сохранения")
            created = updated = 0
            for child_id, text in merged.items():
                cur.execute(f"SELECT id FROM {SCHEMA}.child_weekly_reports WHERE child_id=%s AND week_start=%s ORDER BY id LIMIT 1", (child_id, week_start))
                existing = cur.fetchone()
                if existing:
                    cur.execute(f"UPDATE {SCHEMA}.child_weekly_reports SET report_text=%s, author=COALESCE(%s, author), updated_at=NOW() WHERE id=%s", (text, author, existing["id"]))
                    updated += 1
                else:
                    cur.execute(f"INSERT INTO {SCHEMA}.child_weekly_reports (child_id, author, week_start, report_text) VALUES (%s,%s,%s,%s)", (child_id, author, week_start, text))
                    created += 1
            conn.commit()
            return ok({"created": created, "updated": updated}, 201)

        if method == "GET" and params.get("week_start"):
            return ok({"reports": load_week_reports(cur, monday_of(params["week_start"]))})

        if method == "GET":
            child_id = params.get("child_id")
            if not child_id:
                return err("Параметр child_id обязателен")
            cur.execute(
                f"SELECT * FROM {SCHEMA}.child_weekly_reports WHERE child_id = %s ORDER BY week_start DESC, created_at DESC",
                (child_id,),
            )
            return ok({"reports": [dict(r) for r in cur.fetchall()]})

        if method == "POST":
            child_id = body.get("child_id")
            week_start = body.get("week_start")
            text = (body.get("report_text") or "").strip()
            if not child_id or not week_start:
                return err("Поля child_id и week_start обязательны")
            if not text:
                return err("Введите текст отчёта")
            cur.execute(
                f"INSERT INTO {SCHEMA}.child_weekly_reports (child_id, author, week_start, report_text) VALUES (%s,%s,%s,%s) RETURNING *",
                (child_id, (body.get("author") or "").strip() or None, week_start, text),
            )
            row = cur.fetchone()
            conn.commit()
            return ok({"report": dict(row)}, 201)

        if method == "PUT":
            report_id = params.get("id")
            week_start = body.get("week_start")
            text = (body.get("report_text") or "").strip()
            if not report_id:
                return err("Параметр id обязателен")
            if not week_start or not text:
                return err("Поля week_start и report_text обязательны")
            cur.execute(
                f"UPDATE {SCHEMA}.child_weekly_reports SET week_start=%s, report_text=%s, updated_at=NOW() WHERE id=%s RETURNING *",
                (week_start, text, report_id),
            )
            row = cur.fetchone()
            conn.commit()
            if not row:
                return err("Отчёт не найден", 404)
            return ok({"report": dict(row)})

        if method == "DELETE":
            report_id = params.get("id")
            if not report_id:
                return err("Параметр id обязателен")
            cur.execute(f"DELETE FROM {SCHEMA}.child_weekly_reports WHERE id=%s RETURNING id", (report_id,))
            row = cur.fetchone()
            conn.commit()
            if not row:
                return err("Отчёт не найден", 404)
            return ok({"success": True})

        return err("Метод не поддерживается", 405)
    finally:
        conn.close()
